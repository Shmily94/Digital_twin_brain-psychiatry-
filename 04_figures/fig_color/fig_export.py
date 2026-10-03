"""Editable-text PPTX export (extracted from the project figure skill).
Graphics are rasterised; every matplotlib Text becomes a native PowerPoint text box."""
import os
import matplotlib as mpl
import matplotlib.pyplot as plt

ARIAL_STACK = ['Arial', 'Helvetica', 'Liberation Sans', 'DejaVu Sans']
BOLD_WEIGHTS = ('bold', 'semibold', 'demibold', 'heavy', 'black', 'extra bold')

def collect_text_records(fig):
    fig.canvas.draw()
    rend = fig.canvas.get_renderer()
    recs = []
    for t in fig.findobj(mpl.text.Text):
        s = t.get_text()
        if not s.strip() or not t.get_visible():
            continue
        try:
            bb = t.get_window_extent(rend)
        except Exception:
            continue
        rot = float(t.get_rotation()) % 360.0
        if min(rot % 180.0, 180.0 - (rot % 180.0)) > 0.5:
            t.set_rotation(0)
            bb0 = t.get_window_extent(rend)
            t.set_rotation(rot)
            w, h = bb0.width, bb0.height
        else:
            w, h = bb.width, bb.height
        recs.append(dict(text=s, cx=0.5 * (bb.x0 + bb.x1), cy=0.5 * (bb.y0 + bb.y1),
                         w=w, h=h, rot=rot, size=float(t.get_fontsize()),
                         color=t.get_color(), ha=t.get_ha(),
                         weight=str(t.get_fontweight()).lower(),
                         style=str(t.get_fontstyle()).lower(), obj=t))
    return recs

def fig_to_pptx(fig, pptx_path, dpi=400, keep_background_png=False):
    """One slide: raster graphics layer + one native PowerPoint text box per label."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from matplotlib import colors as mcolors

    recs = collect_text_records(fig)
    fw, fh = [float(v) for v in fig.get_size_inches()]
    fdpi = float(fig.dpi)

    vis = [(r["obj"], r["obj"].get_visible()) for r in recs]
    for obj, _ in vis:
        obj.set_visible(False)
    bg = os.path.splitext(pptx_path)[0] + "_bg.png"
    fig.savefig(bg, dpi=dpi, facecolor=fig.get_facecolor())
    for obj, v in vis:
        obj.set_visible(v)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(fw), Inches(fh)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(bg, 0, 0, width=Inches(fw), height=Inches(fh))

    align = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    pad = 1.25
    for r in recs:
        w, h = r["w"] * pad / fdpi, r["h"] * pad / fdpi
        left = r["cx"] / fdpi - w / 2.0
        top = (fh - r["cy"] / fdpi) - h / 2.0
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = align.get(r["ha"], PP_ALIGN.CENTER)
        run = p.add_run()
        run.text = r["text"].replace("$", "").replace("\\", "")
        f = run.font
        f.name = "Arial"
        f.size = Pt(r["size"])
        f.bold = r["weight"] in BOLD_WEIGHTS or r["weight"].isdigit() and int(r["weight"]) >= 600
        f.italic = r["style"] in ("italic", "oblique")
        rgb = mcolors.to_rgb(r["color"])
        f.color.rgb = RGBColor(*[int(round(255 * c)) for c in rgb])
        if r["rot"]:
            box.rotation = (-r["rot"]) % 360.0
    prs.save(pptx_path)
    if not keep_background_png:
        try:
            os.remove(bg)
        except OSError:
            pass
    return pptx_path

def check_fonts(fig, want="Arial"):
    """Return [(text, resolved_font_file)] for text objects not resolving to `want`."""
    from matplotlib import font_manager as fmg
    bad = []
    for t in fig.findobj(mpl.text.Text):
        if not t.get_text().strip() or not t.get_visible():
            continue
        try:
            path = fmg.findfont(t.get_fontproperties(), fallback_to_default=True)
        except Exception:
            path = "?"
        if want.lower().replace(" ", "") not in os.path.basename(path).lower():
            bad.append((t.get_text()[:40], os.path.basename(path)))
    return bad

