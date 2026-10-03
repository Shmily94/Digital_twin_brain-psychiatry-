"""Render the Fig-1 scene to PNG/PDF (matplotlib) and to an editable PPTX (python-pptx)."""
import numpy as np, matplotlib as mpl, matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Rectangle, Ellipse, FancyArrowPatch, Polygon
from fig1_scene import build, FIG_W, FIG_H

DASH = {None: 'solid', 'dash': (0, (2.6, 1.7)), 'dot': (0, (1.9, 1.3))}

def render_mpl(scene, out_png='fig1_framework.png', out_pdf='fig1_framework.pdf', dpi=450):
    apply_figure_style(frame='none', sizes=(8, 7, 6))
    fig = plt.figure(figsize=(FIG_W, FIG_H))
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, 100); ax.set_ylim(0, 100); ax.axis('off')
    ar = FIG_W / FIG_H
    for o in sorted(scene, key=lambda d: d['z']):
        k = o['k']
        if k == 'rect':
            kw = dict(facecolor=o['fill'] or 'none', edgecolor=o['line'] or 'none',
                      lw=o['lw'], zorder=o['z'])
            if o['round_']:
                ax.add_patch(FancyBboxPatch((o['x'], o['y']), o['w'], o['h'],
                                            boxstyle=f"round,pad=0,rounding_size={o['round_']}", **kw))
            else:
                ax.add_patch(Rectangle((o['x'], o['y']), o['w'], o['h'], **kw))
        elif k == 'oval':
            ax.add_patch(Ellipse((o['cx'], o['cy']), 2 * o['rx'], 2 * o['ry'],
                                 facecolor=o['fill'] or 'none', edgecolor=o['line'] or 'none',
                                 lw=o['lw'], zorder=o['z']))
        elif k == 'line':
            if o['arrow']:
                ax.add_patch(FancyArrowPatch((o['x1'], o['y1']), (o['x2'], o['y2']),
                                             arrowstyle='-|>', mutation_scale=3.6 + 2.4 * o['lw'],
                                             lw=o['lw'], color=o['color'], zorder=o['z'],
                                             linestyle=DASH[o['dash']], shrinkA=0, shrinkB=0))
            else:
                ax.plot([o['x1'], o['x2']], [o['y1'], o['y2']], lw=o['lw'], color=o['color'],
                        ls=DASH[o['dash']], zorder=o['z'], solid_capstyle='round')
        elif k == 'poly':
            ax.add_patch(Polygon(o['pts'], closed=o['close'], facecolor=o['fill'] or 'none',
                                 edgecolor=o['line'] or 'none', lw=o['lw'], zorder=o['z']))
        elif k == 'path':
            p = np.array(o['pts'])
            ax.plot(p[:, 0], p[:, 1], lw=o['lw'], color=o['color'], zorder=o['z'],
                    solid_capstyle='round')
            if o['arrow']:
                ax.add_patch(FancyArrowPatch(p[-3], p[-1], arrowstyle='-|>', mutation_scale=6,
                                             lw=o['lw'], color=o['color'], zorder=o['z'],
                                             shrinkA=0, shrinkB=0))
        elif k == 'text':
            ax.text(o['x'], o['y'], o['s'], fontsize=o['size'], color=o['color'],
                    weight='bold' if o['bold'] else 'normal', ha=o['ha'], va=o['va'],
                    rotation=o['rot'], linespacing=o['ls'], zorder=o['z'],
                    rotation_mode='anchor' if o['rot'] else None)
    fig.savefig(out_png, dpi=dpi); fig.savefig(out_pdf)
    r = fig.canvas.get_renderer()
    tx = [(t, t.get_window_extent(r)) for t in fig.findobj(mpl.text.Text)
          if t.get_text().strip() and t.get_visible()]
    ov = [(p.get_text()[:28], q.get_text()[:28]) for i, (p, bp) in enumerate(tx)
          for q, bq in tx[i + 1:] if bp.overlaps(bq)]
    return fig, len(tx), ov


# ------------------------------------------------------------------ PPTX ----
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
import copy

SCALE = 2.0                  # slide is 2x print size so 5-pt type is editable on screen
FONT = 'Arial'

def _rgb(h): return RGBColor.from_string(h.lstrip('#').upper())

class Pptx:
    def __init__(self, scale=SCALE):
        self.sc = scale
        self.prs = Presentation()
        self.prs.slide_width = Inches(FIG_W * scale)
        self.prs.slide_height = Inches(FIG_H * scale)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.sh = self.slide.shapes
    # coordinate helpers: scene y is bottom-up, PowerPoint is top-down
    def X(self, x): return Emu(int(x / 100 * FIG_W * self.sc * 914400))
    def Y(self, y): return Emu(int((1 - y / 100) * FIG_H * self.sc * 914400))
    def DX(self, w): return Emu(int(w / 100 * FIG_W * self.sc * 914400))
    def DY(self, h): return Emu(int(h / 100 * FIG_H * self.sc * 914400))
    def LW(self, lw): return Pt(lw * self.sc)

    def _style(self, shp, fill, line, lw):
        if fill: shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(fill)
        else: shp.fill.background()
        if line:
            shp.line.color.rgb = _rgb(line); shp.line.width = self.LW(lw)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False

    def rect(self, o):
        if o['round_']:
            s = self.sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, self.X(o['x']), self.Y(o['y'] + o['h']),
                                  self.DX(o['w']), self.DY(o['h']))
            s.adjustments[0] = min(0.5, o['round_'] / min(o['w'], o['h']) * 0.5)
        else:
            s = self.sh.add_shape(MSO_SHAPE.RECTANGLE, self.X(o['x']), self.Y(o['y'] + o['h']),
                                  self.DX(o['w']), self.DY(o['h']))
        self._style(s, o['fill'], o['line'], o['lw'])
        return s

    def oval(self, o):
        s = self.sh.add_shape(MSO_SHAPE.OVAL, self.X(o['cx'] - o['rx']), self.Y(o['cy'] + o['ry']),
                              self.DX(2 * o['rx']), self.DY(2 * o['ry']))
        self._style(s, o['fill'], o['line'], o['lw'])
        return s

    def _arrowhead(self, shp, size='med'):
        ln = shp.line._get_or_add_ln()
        he = ln.makeelement(qn('a:headEnd'), {})
        te = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'sm', 'len': 'sm'})
        ln.append(he); ln.append(te)

    def line(self, o):
        s = self.sh.add_connector(MSO_CONNECTOR.STRAIGHT, self.X(o['x1']), self.Y(o['y1']),
                                  self.X(o['x2']), self.Y(o['y2']))
        s.line.color.rgb = _rgb(o['color']); s.line.width = self.LW(o['lw'])
        if o['dash'] == 'dash': s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        elif o['dash'] == 'dot': s.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
        if o['arrow']: self._arrowhead(s)
        return s

    def _freeform(self, pts, fill, line, lw, close):
        b = self.sh.build_freeform(self.X(pts[0][0]), self.Y(pts[0][1]))
        b.add_line_segments([(self.X(x), self.Y(y)) for x, y in pts[1:]], close=close)
        s = b.convert_to_shape()
        self._style(s, fill, line, lw)
        return s

    def poly(self, o):
        return self._freeform(o['pts'], o['fill'], o['line'], o['lw'], o['close'])

    def path(self, o):
        s = self._freeform(o['pts'], None, o['color'], o['lw'], False)
        if o['arrow']: self._arrowhead(s)
        return s

    def text(self, o):
        nlines = o['s'].count('\n') + 1
        pt = o['size'] * self.sc
        hgt = Emu(int(nlines * pt * o['ls'] * 1.18 * 12700))
        wid = Emu(int(max(8, max(len(l) for l in o['s'].split('\n'))) * pt * 0.62 * 12700))
        if o['rot'] in (90, -90):
            cx, cy = self.X(o['x']), self.Y(o['y'])
            left, top = Emu(int(cx - wid / 2)), Emu(int(cy - hgt / 2))
        else:
            if o['ha'] == 'left':   left = self.X(o['x'])
            elif o['ha'] == 'center': left = Emu(int(self.X(o['x']) - wid / 2))
            else:                   left = Emu(int(self.X(o['x']) - wid))
            if o['va'] == 'top':      top = self.Y(o['y'])
            elif o['va'] == 'center': top = Emu(int(self.Y(o['y']) - hgt / 2))
            else:                     top = Emu(int(self.Y(o['y']) - hgt))
        tb = self.sh.add_textbox(left, top, wid, hgt)
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'center': MSO_ANCHOR.MIDDLE,
                              'bottom': MSO_ANCHOR.BOTTOM}[o['va']]
        for i, ln in enumerate(o['s'].split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
                           'right': PP_ALIGN.RIGHT}[o['ha']]
            p.line_spacing = o['ls']
            r = p.add_run(); r.text = ln
            r.font.size = Pt(pt); r.font.bold = o['bold']; r.font.name = FONT
            r.font.color.rgb = _rgb(o['color'])
        if o['rot']:
            tb.rotation = -o['rot'] % 360
        return tb


def render_pptx(scene, out='fig1_framework_editable.pptx', scale=SCALE):
    p = Pptx(scale)
    disp = dict(rect=p.rect, oval=p.oval, line=p.line, poly=p.poly, path=p.path, text=p.text)
    import re as _re
    n = 0
    for i, o in enumerate(sorted(scene, key=lambda d: d['z'])):
        s = disp[o['k']](o)
        if o['k'] == 'text':
            slug = _re.sub(r'[^A-Za-z0-9]+', '_', o['s'].splitlines()[0])[:22].strip('_')
            s.name = f'txt{i:03d}_{slug or "t"}'
        else:
            s.name = f'{o["k"]}{i:03d}'
        n += 1
    p.prs.save(out)
    return out, n
