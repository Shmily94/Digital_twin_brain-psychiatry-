"""Build an editable PowerPoint version of the Fig. 1 PIPELINE layout.

Companion to fig1_ppt.py (which does the same for the a-f panel layout).
Geometry mirrors fig1_pipeline.py 1:1 on a 180 x 150 mm slide, so the deck opens
looking like fig1_pipeline.png; every element is then movable on its own:

  * stage cards, header bands, number badges, colour swatches  -> PowerPoint shapes
  * all prose, sample-selection counts and plot labels         -> native Arial text
  * each stage's data graphic                                  -> its own transparent
    PNG at final size, 600 dpi, with a matching PDF in fig1_ppt_parts/

Nothing here is schematic: every plotted value comes from fig1_data/*.csv and
fig.2/fig2_data/fig2c_profile_scores.csv, the same extracts fig1_pipeline.py uses.

Outputs: fig1_pipeline_editable.pptx, fig1_ppt_parts/pipe_*.png|pdf
"""
import os, sys, textwrap
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.colors as mcolors
from matplotlib.patches import FancyBboxPatch, Circle
from PIL import Image
from pptx import Presentation
from pptx.util import Mm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, '..', 'fig_color'))
from np_dtb_style import apply_np_style, C, LW        # noqa: E402
from supp_kit import fill, pt_edge                    # noqa: E402
apply_np_style()

D = os.path.join(HERE, 'fig1_data')
P = os.path.join(HERE, 'fig1_ppt_parts')
F2 = os.path.join(HERE, '..', 'fig.2', 'fig2_data')
os.makedirs(P, exist_ok=True)

INK, SOFT, FAINT, HAIR = '#1A1A1A', '#4D4D4D', '#8C8C8C', '#D9D9D9'
BAND, RULE = '#F0F0F0', '#E0E0E0'

SEL = pd.read_csv(f'{D}/fig1_sample_selection.csv')
FID = pd.read_csv(f'{D}/fig1_mini_fidelity.csv').sort_values('bold_r_mean', ascending=False)
PER = pd.read_csv(f'{D}/fig1_mini_perturbation_n288.csv')
H27 = pd.read_csv(f'{D}/fig1_mini_healthy_n27.csv')
M36 = pd.read_csv(f'{D}/fig1_mini_clinical_n36.csv')
L85 = pd.read_csv(f'{D}/fig1_mini_longitudinal_n85.csv')
EDG = pd.read_csv(f'{D}/fig1_mini_np12_edges.csv')
STR = pd.read_csv(f'{F2}/fig2c_profile_scores.csv')

# ------------------------------------------------------------ part export ----
def part(name, w_mm, h_mm, draw, dpi=600):
    """render one graphic to fig1_ppt_parts/<name>.png|pdf on a transparent ground"""
    fig = plt.figure(figsize=(w_mm / 25.4, h_mm / 25.4))
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis('off')
    draw(ax)
    fig.savefig(f'{P}/{name}.png', dpi=dpi, transparent=True)
    fig.savefig(f'{P}/{name}.pdf', transparent=True)
    plt.close(fig)
    return f'{P}/{name}.png'


def bx(ax, xc, v, col, ylim, w=.30, lwf=1.0):
    """box-and-whisker in axes-fraction coordinates; ylim in data units"""
    v = np.asarray(v, float)
    lo_, hi_ = ylim
    Y = lambda t: (np.asarray(t, float) - lo_) / (hi_ - lo_)
    q1, med, q3 = np.percentile(v, [25, 50, 75]); iqr = q3 - q1
    lo = v[v >= q1 - 1.5 * iqr].min(); hi = v[v <= q3 + 1.5 * iqr].max()
    ax.plot([xc, xc], [Y(lo), Y(hi)], color=INK, lw=LW * .7 * lwf, zorder=3)
    ax.add_patch(FancyBboxPatch((xc - w / 2, Y(q1)), w, Y(q3) - Y(q1),
                                boxstyle='round,pad=0,rounding_size=0.004',
                                facecolor=fill(col), edgecolor=INK,
                                lw=LW * .7 * lwf, zorder=4))
    ax.plot([xc - w / 2, xc + w / 2], [Y(med)] * 2, color=INK, lw=LW * .9 * lwf, zorder=5)


# ---- stage 1: NP ring, and the STRATIFY boxes ------------------------------
def _ring(ax):
    nodes = sorted(pd.unique(EDG[['i_217', 'j_217']].values.ravel()))
    th = np.linspace(90, -270, len(nodes), endpoint=False) * np.pi / 180
    R = .45
    pos = {n: (.5 + R * np.cos(t), .5 + R * np.sin(t)) for n, t in zip(nodes, th)}
    for _, e in EDG.iterrows():
        a_, b_ = pos[e.i_217], pos[e.j_217]
        ax.plot([a_[0], b_[0]], [a_[1], b_[1]], color=C('np12'), lw=.9, alpha=.75,
                zorder=3, solid_capstyle='round')
    for n, (px, py) in pos.items():
        ax.add_patch(Circle((px, py), .042, facecolor='white', edgecolor=C('np12'),
                            linewidth=LW * .8, zorder=4))


def _strat(ax):
    yl = (-4.5, 4.5)
    bx(ax, .28, STR.loc[STR.Group == 'HC', 'NP_factor'].values, C('hc'), yl, w=.30)
    bx(ax, .74, STR.loc[STR.Group == 'Patient', 'NP_factor'].values, C('patient'), yl, w=.30)


# ---- stage 2: assimilated-region BOLD r, seven builds ----------------------
LBL = {'10m_268': '10 M/268', '3m_268': '3 M/268', '1b': '1 B', '100m': '100 M',
       '10m': '10 M vox', '10m_own': '10 M own', '10m_1000': '10 M/1000'}
FAM = {'10m_268': 'model_regional', '3m_268': 'model_regional', '10m_1000': 'model_regional',
       '1b': 'model_voxel', '100m': 'model_voxel', '10m': 'model_voxel', '10m_own': 'model_voxel'}


def _bars(ax):
    n = len(FID)
    lo_, hi_ = .6, .95
    for k, (_, r) in enumerate(FID.iterrows()):
        xc = (k + .5) / n
        h = (r.bold_r_mean - lo_) / (hi_ - lo_)
        ax.add_patch(FancyBboxPatch((xc - .34 / n, 0), .68 / n, h,
                                    boxstyle='round,pad=0,rounding_size=0.004',
                                    facecolor=C(FAM[r.model]), edgecolor='none', zorder=3))


# ---- stage 3: baseline -> AMPA -> GABA-A ----------------------------------
def _perturb(ax):
    yl = (-3, 8)
    for k, (c, col) in enumerate([('simulated', 'baseline'), ('ampa', 'ampa'), ('gaba', 'gaba')]):
        bx(ax, (k + .5) / 3, PER[c].values, C(col), yl, w=.34 / 3 * 2)


# ---- stage 4: the two pharmacological datasets ----------------------------
def _healthy(ax):
    yl = (-1.2, 2.3)
    for k, (c, col) in enumerate([('Placebo', 'placebo'), ('Ketamine', 'ketamine'),
                                  ('Midazolam', 'midazolam')]):
        bx(ax, (k + .5) / 3, H27[c].values, C(col), yl, w=.26)


def _clinical(ax):
    yl = (-4, 4)
    for k, (c, col) in enumerate([('FC_p2', 'placebo'), ('FC_d2', 'ketamine')]):
        bx(ax, (k + .5) / 2, M36[c].values, C(col), yl, w=.34)


# ---- stage 5: AMPA restoration index vs four-year symptom change ----------
XV, YV = L85.ampa_restoration_index.values, L85.fu3_symptom_change.values


def _scatter(ax):
    xl = (XV.min() - .1 * np.ptp(XV), XV.max() + .1 * np.ptp(XV))
    yl = (YV.min() - .1 * np.ptp(YV), YV.max() + .1 * np.ptp(YV))
    X = lambda v: (np.asarray(v, float) - xl[0]) / (xl[1] - xl[0])
    Y = lambda v: (np.asarray(v, float) - yl[0]) / (yl[1] - yl[0])
    ax.scatter(X(XV), Y(YV), s=2.2, facecolor=C('ampa'), edgecolor=pt_edge(C('ampa')),
               linewidth=LW * .3, zorder=3)
    b, a = np.polyfit(XV, YV, 1)
    xx = np.array([XV.min(), XV.max()])
    ax.plot(X(xx), Y(a + b * xx), color=C('ampa'), lw=LW * 1.5, zorder=4)


# ------------------------------------------------------------------ deck -----
SW, SH = 180.0, 150.0
prs = Presentation()
prs.slide_width, prs.slide_height = Mm(SW), Mm(SH)
sl = prs.slides.add_slide(prs.slide_layouts[6])
UX, UY = SW / 100.0, SH / 100.0
AR = SW / SH


def mx(u):
    return Mm(u * UX)


def my(u):                      # canvas y (0 = bottom) -> slide y (0 = top)
    return Mm((100 - u) * UY)


def rgb(h):
    return RGBColor.from_string(mcolors.to_hex(h).lstrip('#').upper())


def shape(kind, x, y, w, h, fc=None, lc=None, lw=.6, txt=None, size=5.0, col=INK,
          bold=False, align=PP_ALIGN.CENTER, radius=None):
    """x, y = canvas units of the TOP-LEFT corner; w, h = canvas units"""
    s = sl.shapes.add_shape(kind, mx(x), my(y), Mm(w * UX), Mm(h * UY))
    if fc is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fc)
    if lc is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(lc); s.line.width = Pt(lw)
    s.shadow.inherit = False
    if radius is not None and s.adjustments:
        s.adjustments[0] = radius
    tf = s.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = txt or ''
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(col)
    r.font.name = 'Arial'
    return s


def text(x, y, s_, size=4.6, col=SOFT, bold=False, w=40, h=4, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, spacing=1.25):
    """x, y = canvas units of the box's TOP-LEFT corner; w, h in canvas units"""
    tb = sl.shapes.add_textbox(mx(x), my(y), Mm(w * UX), Mm(h * UY))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(s_).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = rgb(col); r.font.name = 'Arial'
    return tb


def rich(x, y, runs, size=4.9, col=INK, w=20, h=2.4, spacing=1.25):
    """one line, several runs: [(text, bold), ...] -- for '1,050 analysed'"""
    tb = sl.shapes.add_textbox(mx(x), my(y), Mm(w * UX), Mm(h * UY))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.line_spacing = spacing
    for s_, b in runs:
        r = p.add_run(); r.text = s_
        r.font.size = Pt(size); r.font.bold = b
        r.font.color.rgb = rgb(col); r.font.name = 'Arial'
    return tb


def pic(path, x, y, w):
    """place an image with its aspect preserved; x, y = top-left in canvas units"""
    im = Image.open(path)
    h = w * (im.size[1] / im.size[0]) * (UX / UY)
    sl.shapes.add_picture(path, mx(x), my(y), width=Mm(w * UX))
    return h


def arrow(kind, x, y, w, h, col):
    s = sl.shapes.add_shape(kind, mx(x), my(y), Mm(w * UX), Mm(h * UY))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(col)
    s.line.fill.background(); s.shadow.inherit = False
    return s


def line(x1, y1, x2, y2, col=RULE, lw=.6):
    from pptx.util import Emu
    s = sl.shapes.add_connector(1, mx(x1), my(y1), mx(x2), my(y2))
    s.line.color.rgb = rgb(col); s.line.width = Pt(lw)
    return s


RECT, RRECT, OVAL = MSO_SHAPE.RECTANGLE, MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.OVAL

# --- title -------------------------------------------------------------------
text(2.0, 99.0, 'Fig. 1 | Analytical pipeline, from sample selection to final outputs',
     size=8.4, col=INK, bold=True, w=96, h=4)
text(2.0, 95.0, 'A symptom-linked circuit is derived empirically, rebuilt inside individual digital '
                'twin brains and perturbed in silico; the simulated response is then tested\n'
                'against two pharmacological datasets and against four-year symptom change. '
                'The NP factor is the single quantity carried through all five stages.',
     size=5.6, col=SOFT, w=96, h=6, spacing=1.45)

X0, GAP = 2.0, 1.15
CW = (100 - 2 * X0 - 4 * GAP) / 5
CY, CH = 18.5, 64.0
xs = [X0 + i * (CW + GAP) for i in range(5)]

TITLES = ['Derive the circuit\nphenotype', 'Build and select\nthe digital twin',
          'Perturb the circuit\nin silico', 'Validate against\ntwo drug datasets',
          'Forecast four-year\nsymptom change']
TAKE = ['The 29-edge negative profile,\nand its 12-edge NP factor,\nseparate patients from controls',
        'Build choice decides which\ninference is admissible \u2014\nit is not a free parameter',
        'Both perturbations move NP\nconnectivity, in a\nparticipant-specific direction',
        'Drug effects follow the\nvirtual perturbation without\nany refitting of the mapping',
        'The perturbational response\ncarries prospective\ninformation about symptoms']
FIGREF = ['Fig. 2 \u00b7 Supp. STRATIFY, PET', 'Fig. 3 \u00b7 Supp. assimilation region',
          'Fig. 4 \u00b7 Supp. conductance grid',
          'Fig. 5a\u2013g \u00b7 Supp. paired MID,\nwhole-brain drug, Oldham',
          'Fig. 5h \u00b7 Supp. longitudinal']

for i, x in enumerate(xs):
    shape(RRECT, x, CY + CH, CW, CH, fc='white', lc=HAIR, lw=.6, radius=.035)
    shape(RECT, x, CY + CH, CW, 7.4, fc=BAND)
    shape(OVAL, x + 1.2, CY + CH - 2.5, 2.4, 2.4 * AR, fc=SOFT, txt=str(i + 1),
          size=5.8, col='white', bold=True)
    text(x + 4.6, CY + CH - 1.6, TITLES[i], size=6.4, col=INK, bold=True,
         w=CW - 5.0, h=5.0, spacing=1.2)
    text(x + .9, CY + 11.0, TAKE[i], size=5.0, col=SOFT, w=CW - 1.8, h=6.0,
         spacing=1.45, anchor=MSO_ANCHOR.MIDDLE)
    text(x + .9, CY + 4.6, FIGREF[i], size=4.6, col=FAINT, italic=True,
         w=CW - 1.8, h=4.0)
    line(x + .9, CY + 11.8, x + CW - .9, CY + 11.8)
    if i:
        arrow(MSO_SHAPE.RIGHT_ARROW, x - GAP + .10, CY + CH / 2 + .55, 1.1, 1.1, '#B0B0B0')

# --- sample-selection strand -------------------------------------------------
SY = CY + CH - 20.0
for i, x in enumerate(xs):
    text(x + .9, CY + CH - 8.6, 'SAMPLE SELECTION', size=4.5, col=FAINT, bold=True,
         w=CW - 1.8, h=2.2)
    yy = SY + 7.6
    for _, r in SEL[SEL.stage == i + 1].iterrows():
        text(x + .9, yy + 1.2, r.cohort, size=5.0, col=INK, bold=True, w=CW - 1.8, h=2.4)
        if r.excluded:
            text(x + .9, yy - 1.6, f'{r.screened:,} screened', size=4.7, col=SOFT,
                 w=CW - 1.8, h=2.2)
            wrapped = textwrap.fill(f'\u2212{r.excluded:,}  {r.reason_short}', 30)
            nl = wrapped.count('\n')
            text(x + .9, yy - 3.6 - .9 * nl, wrapped, size=4.1, col=C('patient'),
                 w=CW - 1.8, h=2.2 + 1.9 * nl, spacing=1.28)
            rich(x + .9, yy - 6.3 - 1.9 * nl, [(f'{r.analysed:,}', True), (' analysed', False)],
                 size=4.9, col=INK, w=CW - 1.8)
            yy -= 10.4 + 1.9 * nl
        else:
            rich(x + .9, yy - 1.7, [(f'{r.analysed:,}', True), (' analysed', False)],
                 size=4.9, col=INK, w=CW - 1.8)
            yy -= 5.8

# --- stage 1 graphic ---------------------------------------------------------
g1x, g1y, g1w, g1h = xs[0] + 1.0, CY + 17.0, CW - 2.0, 16.0
RING_W = g1w * .52
pic(part('pipe_s1_ring', RING_W * UX, RING_W * UX, _ring), g1x, g1y + g1h - .8, RING_W)
text(g1x, g1y + 1.9, '12-edge NP factor', size=4.5, col=SOFT, w=RING_W + 2, h=2.2,
     align=PP_ALIGN.CENTER)
BW, BH = g1w * .32, g1h * .62
pic(part('pipe_s1_box', BW * UX, BH * UY, _strat), g1x + g1w * .62, g1y + 3.2 + BH, BW)
text(g1x + g1w * .62 + BW * .28 - 2.0, g1y + 2.2, 'HC', size=4.4, col=SOFT, w=4.0, h=2.0,
     align=PP_ALIGN.CENTER)
text(g1x + g1w * .62 + BW * .74 - 2.0, g1y + 2.2, 'Patient', size=4.4, col=SOFT, w=4.0, h=2.0,
     align=PP_ALIGN.CENTER)
text(g1x + g1w * .62, g1y + g1h + .6, 'STRATIFY', size=4.5, col=SOFT, w=BW, h=2.2,
     align=PP_ALIGN.CENTER)

# --- stage 2 graphic ---------------------------------------------------------
g2x, g2y, g2w, g2h = xs[1] + 1.4, CY + 18.5, CW - 3.2, 15.0
B2H = g2h * .74
text(g2x, g2y + B2H + 3.2, 'Assimilated-region BOLD r', size=4.5, col=SOFT, w=g2w, h=2.2)
pic(part('pipe_s2_bars', g2w * UX, B2H * UY, _bars), g2x, g2y + B2H, g2w)
for k, (_, r) in enumerate(FID.iterrows()):
    xc = g2x + (k + .5) / len(FID) * g2w
    text(xc - 2.0, g2y - (0.4 if k % 2 == 0 else 2.0), LBL[r.model], size=3.9, col=SOFT,
         w=4.0, h=1.8, align=PP_ALIGN.CENTER)
for lab, key, dy in [('regional', 'model_regional', 0), ('voxel', 'model_voxel', -1.9)]:
    shape(RECT, g2x + g2w * .60, g2y + B2H + 1.0 + dy, 1.0, .9 * AR, fc=C(key))
    text(g2x + g2w * .60 + 1.4, g2y + B2H + 1.1 + dy, lab, size=4.2, col=SOFT, w=8, h=1.8)
text(g2x, g2y - 3.4, '1 B fidelity reference\n100 M parameter search\n3 M/268 population runs',
     size=4.3, col=FAINT, w=g2w, h=6.0, spacing=1.4)

# --- stage 3 graphic ---------------------------------------------------------
g3x, g3y, g3w, g3h = xs[2] + 2.0, CY + 17.5, CW - 4.0, 16.0
B3H = g3h * .72
text(g3x, g3y + 2.6 + B3H + 2.8, 'Simulated NP factor, n = 288', size=4.5, col=SOFT,
     w=g3w + 2, h=2.2)
pic(part('pipe_s3_box', g3w * UX, B3H * UY, _perturb), g3x, g3y + 2.6 + B3H, g3w)
for k, lab in enumerate(['baseline', 'AMPA \u2191', 'GABA-A \u2191']):
    text(g3x + (k + .5) / 3 * g3w - 2.6, g3y + 2.0, lab, size=4.3, col=SOFT, w=5.2, h=1.8,
         align=PP_ALIGN.CENTER)
text(g3x, g3y - 2.6, 'GABA-A is applied on top of AMPA', size=4.3, col=FAINT, w=g3w + 2, h=2.2)

# --- stage 4 graphic ---------------------------------------------------------
g4x, g4y, g4w, g4h = xs[3] + 1.6, CY + 15.5, CW - 3.4, 17.5
W4A, W4B, H4 = g4w * .52, g4w * .34, 7.4
text(g4x, g4y + 9.0 + H4 + 2.6, 'Healthy, n = 27', size=4.4, col=SOFT, w=W4A + 4, h=2.2)
pic(part('pipe_s4_healthy', W4A * UX, H4 * UY, _healthy), g4x, g4y + 9.0 + H4, W4A)
text(g4x, g4y + 8.6, 'pla \u00b7 ket \u00b7 mid', size=4.0, col=FAINT, w=W4A + 2, h=1.8)
text(g4x + g4w * .62, g4y + 9.0 + H4 + 2.6, 'Clinical, n = 36', size=4.4, col=SOFT,
     w=W4B + 6, h=2.2)
pic(part('pipe_s4_clinical', W4B * UX, H4 * UY, _clinical), g4x + g4w * .62, g4y + 9.0 + H4, W4B)
text(g4x + g4w * .62, g4y + 8.6, 'pla \u00b7 ket', size=4.0, col=FAINT, w=W4B + 2, h=1.8)
text(g4x, g4y + 4.6, 'Summed NP-related MID FC;\nwhole-brain maps tested for\n'
                     'directional concordance', size=4.3, col=FAINT, w=g4w + 2, h=6.0,
     spacing=1.4)

# --- stage 5 graphic ---------------------------------------------------------
g5x, g5y, g5w, g5h = xs[4] + 2.2, CY + 18.0, CW - 4.4, 15.5
B5H = g5h * .80
text(g5x, g5y + 1.4 + B5H + 2.8, 'AMPA index vs four-year\nsymptom change', size=4.4,
     col=SOFT, w=g5w + 2, h=4.0, spacing=1.3)
pic(part('pipe_s5_scatter', g5w * UX, B5H * UY, _scatter), g5x, g5y + 1.4 + B5H, g5w)
text(g5x, g5y - .4, 'n = 85 \u00b7 r = 0.26 \u00b7 model R\u00b2 = 0.23', size=4.3, col=FAINT,
     w=g5w + 2, h=2.2)

# --- linking arrows ----------------------------------------------------------
arrow(MSO_SHAPE.CURVED_UP_ARROW, xs[1] + CW * .5, CY + CH + 3.6, CW + GAP, 2.6,
      C('model_voxel'))
text(xs[1] + CW * .5 - 6.0, CY + CH + 7.4, 'selected build fixes which inference is admissible',
     size=4.7, col=C('model_voxel'), w=CW + GAP + 12, h=2.2, align=PP_ALIGN.CENTER)
arrow(MSO_SHAPE.CURVED_DOWN_ARROW, xs[2] + CW * .5, CY - .6, CW + GAP, 2.6, C('ampa'))
text(xs[2] + CW * .5 - 8.0, CY - 4.0,
     'perturbation \u2192 response mapping applied to empirical data without refitting',
     size=4.7, col=C('ampa'), w=CW + GAP + 16, h=2.2, align=PP_ALIGN.CENTER)

# --- colour correspondence ---------------------------------------------------
text(2.0, 10.8, 'Colour correspondence', size=5.4, col=INK, bold=True, w=40, h=2.6)
KEYS = [('np12', 'NP factor / negative profile'), ('pos_profile', 'positive profile'),
        ('model_regional', 'regional build'), ('model_voxel', 'voxel build'),
        ('ampa', 'virtual AMPA \u2194 ketamine'), ('gaba', 'virtual GABA-A \u2194 midazolam'),
        ('hc', 'HC / baseline / placebo'), ('high_symptom', 'high-symptom'),
        ('patient', 'patient / MDD'), ('aud', 'AUD')]
for k, (key, lab) in enumerate(KEYS):
    col, row = k % 5, k // 5
    x = 2.0 + col * 19.4
    y = 7.2 - row * 3.4
    shape(RECT, x, y, 1.5, 1.5 * AR, fc=C(key))
    text(x + 2.1, y - .1, lab, size=4.6, col=SOFT, w=17, h=2.0)

out = os.path.join(HERE, 'fig1_pipeline_editable.pptx')
prs.save(out)

# ------------------------------------------------------------ verification ---
chk = Presentation(out).slides[0]
oob = []
for s in chk.shapes:
    l, t = s.left / 36000.0, s.top / 36000.0
    w, h = (s.width or 0) / 36000.0, (s.height or 0) / 36000.0
    if l < -0.2 or t < -0.2 or l + w > SW + 0.2 or t + h > SH + 0.2:
        oob.append((s.shape_type, s.name, round(l, 1), round(t, 1), round(w, 1), round(h, 1)))
print(f'wrote {out}')
print(f'objects: {len(chk.shapes.__iter__.__self__._spTree) - 2}; '
      f'parts: {len([f for f in os.listdir(P) if f.startswith("pipe_") and f.endswith(".png")])}; '
      f'out of bounds: {len(oob)}')
for o in oob[:8]:
    print('  OOB', o)
